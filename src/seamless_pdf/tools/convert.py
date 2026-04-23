"""Convert to PDF (office, images, text, Markdown, HTML, EPUB, comics, web
pages) and convert from PDF (Excel, PowerPoint, OpenDocument sheets and
slides, EPUB, HTML, Markdown, text, images).

PDF to Word / ODT / RTF are deliberately not offered: every free converter
of good quality is AGPL (pdf2docx/PyMuPDF)."""
from __future__ import annotations

import html as _html
import io
import os
import subprocess
import tempfile
import zipfile

import pikepdf

from ..engine import EditError
from .base import FILES, NONE, PAGES_OPT, PAPER_ENUM, PDF, find_program, paper, parse_pages, tool
from .render import (chrome_print, html_to_pdf, markdown_to_html, open_bytes, soffice_convert,
                     text_to_html)

TO, FROM = "Convert to PDF", "Convert from PDF"

OFFICE = {"doc", "docx", "docm", "dot", "dotx", "ppt", "pptx", "pps", "ppsx", "xls", "xlsx", "xlsm",
          "odt", "ods", "odp", "odg", "rtf", "pub", "vsd", "vsdx", "wpd", "csv", "pages", "numbers",
          "key", "wps", "et", "dps", "fodt", "fods", "fodp"}
IMAGES = {"jpg", "jpeg", "png", "webp", "gif", "bmp", "tif", "tiff", "heic", "heif", "jfif",
          "pnm", "pbm", "pgm", "ppm", "jp2", "jpx", "psd"}


# -- images -----------------------------------------------------------------------
def _image_frames(data: bytes, ext: str) -> list[tuple[bytes, int, int]]:
    """Image file -> list of (png/jpeg bytes, width px, height px) frames."""
    from PIL import Image, ImageOps
    if ext in ("heic", "heif"):
        import pillow_heif
        pillow_heif.register_heif_opener()
    try:
        im = Image.open(io.BytesIO(data))
    except Exception as e:
        raise EditError(f"cannot read image: {e}")
    frames = []
    for i in range(getattr(im, "n_frames", 1)):
        im.seek(i)
        fr = ImageOps.exif_transpose(im.copy())
        buf = io.BytesIO()
        if fr.mode in ("RGBA", "LA", "P", "PA"):
            fr = fr.convert("RGBA")
            fr.save(buf, "PNG")
        else:
            fr = fr.convert("RGB")
            fr.save(buf, "JPEG", quality=92)
        frames.append((buf.getvalue(), fr.width, fr.height))
        if ext not in ("tif", "tiff", "gif"):
            break
    return frames


def _images_pdf(items: list[tuple[bytes, int, int]], page_size="fit", orientation="auto",
                margin=0.0) -> pikepdf.Pdf:
    from ..engine.images import image_xobject
    out = pikepdf.new()
    for data, w, h in items:
        if page_size == "fit":
            pw, ph = w * 72 / 96 + 2 * margin, h * 72 / 96 + 2 * margin
        else:
            orient = orientation if orientation != "auto" else ("landscape" if w > h else "portrait")
            pw, ph = paper(page_size, orient)
        page = out.add_blank_page(page_size=(pw, ph))
        stream, iw, ih = image_xobject(out, data)
        bw, bh = pw - 2 * margin, ph - 2 * margin
        s = min(bw / iw, bh / ih)
        dw, dh = iw * s, ih * s
        x0, y0 = margin + (bw - dw) / 2, margin + (bh - dh) / 2
        name = page.add_resource(out.make_indirect(stream), pikepdf.Name.XObject, prefix="Im")
        page.contents_add(out.make_stream(f"q {dw:.3f} 0 0 {dh:.3f} {x0:.3f} {y0:.3f} cm {name} Do Q".encode()))
    return out


def _svg_pdf(data: bytes, name: str = "image.svg") -> bytes:
    """Vector SVG -> one-page PDF sized to the drawing (Chrome)."""
    import re
    head = data[:4000].decode("utf-8", "ignore")

    def dim(attr):
        m = re.search(rf'<svg[^>]*?\s{attr}\s*=\s*["\']([\d.]+)\s*(px|pt|mm|cm|in)?["\']', head)
        if not m:
            return None
        v, u = float(m.group(1)), m.group(2) or "px"
        return v * {"px": 0.75, "pt": 1, "mm": 72 / 25.4, "cm": 72 / 2.54, "in": 72}[u]
    w, h = dim("width"), dim("height")
    if not (w and h):
        vb = re.search(r'viewBox\s*=\s*["\'][\d.\-]+[ ,]+[\d.\-]+[ ,]+([\d.]+)[ ,]+([\d.]+)["\']', head)
        w, h = (float(vb.group(1)) * 0.75, float(vb.group(2)) * 0.75) if vb else (595, 842)
    if not find_program("chrome"):
        return soffice_convert(data, name, "pdf")
    with tempfile.TemporaryDirectory(prefix="svg-") as tmp:
        with open(os.path.join(tmp, "img.svg"), "wb") as fh:
            fh.write(data)
        path = os.path.join(tmp, "doc.html")
        with open(path, "w") as fh:
            fh.write(f"<html><head><style>@page{{size:{w:.2f}pt {h:.2f}pt;margin:0}}html,body{{margin:0}}"
                     f"img{{display:block;width:{w:.2f}pt;height:{h:.2f}pt}}</style></head>"
                     f"<body><img src='img.svg'></body></html>")
        pdf = open_bytes(chrome_print("file://" + path))
        while len(pdf.pages) > 1:
            del pdf.pages[-1]
        buf = io.BytesIO()
        pdf.save(buf)
        return buf.getvalue()


@tool("images_to_pdf", "Images to PDF", TO,
      "Combine images (JPG, PNG, WEBP, HEIC, TIFF, GIF, BMP, SVG ...) into one "
      "PDF, one image per page.",
      inputs=FILES, accept="image/*,.heic,.heif,.svg",
      aliases=["jpg_to_pdf", "png_to_pdf", "webp_to_pdf", "heic_to_pdf", "tiff_to_pdf",
               "photos_to_pdf", "image_to_pdf"],
      options={"page_size": {"type": "string", "enum": ["fit"] + PAPER_ENUM, "default": "fit"},
               "orientation": {"type": "string", "enum": ["auto", "portrait", "landscape"], "default": "auto"},
               "margin": {"type": "number", "default": 0},
               "output_name": {"type": "string"}})
def images_to_pdf(ctx, page_size="fit", orientation="auto", margin=0, output_name=None):
    from .render import page_form, place_form
    out = pikepdf.new()
    for inp in ctx.inputs:
        if inp.ext == "svg":
            svg = open_bytes(_svg_pdf(inp.data, inp.name))
            if page_size == "fit":
                out.pages.append(svg.pages[0])
            else:
                from .render import page_size as psize
                w, h = psize(svg.pages[0])
                orient = orientation if orientation != "auto" else ("landscape" if w > h else "portrait")
                pw, ph = paper(page_size, orient)
                sheet = out.add_blank_page(page_size=(pw, ph))
                place_form(out, sheet, page_form(out, svg, 0), [margin, margin, pw - margin, ph - margin])
        else:
            out.pages.extend(_images_pdf(_image_frames(inp.data, inp.ext), page_size, orientation, margin).pages)
    ctx.add_pdf(out, output_name or (ctx.inputs[0].stem if len(ctx.inputs) == 1 else "images.pdf"))
    ctx.result = {"pages": len(out.pages)}


def _epub_pdf(data: bytes, size: str) -> bytes:
    exe = find_program("pandoc")
    if not exe:
        raise EditError("pandoc is needed to read EPUB files (brew install pandoc)")
    with tempfile.TemporaryDirectory(prefix="epub-") as tmp:
        src = os.path.join(tmp, "book.epub")
        with open(src, "wb") as fh:
            fh.write(data)
        r = subprocess.run([exe, src, "-f", "epub", "-t", "html5", "--standalone", "--embed-resources",
                            "--metadata", "pagetitle=book"], capture_output=True, timeout=180, cwd=tmp)
        if r.returncode != 0:
            raise EditError(f"pandoc could not read the EPUB: {r.stderr.decode(errors='ignore')[-300:]}")
        return html_to_pdf(r.stdout.decode("utf-8", "replace"), size=size)


def _cbz_pdf(data: bytes) -> bytes:
    items = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for n in sorted(z.namelist()):
            ext = n.rsplit(".", 1)[-1].lower()
            if ext in IMAGES and not n.startswith("__MACOSX"):
                items.extend(_image_frames(z.read(n), ext))
    if not items:
        raise EditError("no images found in the comic archive")
    buf = io.BytesIO()
    _images_pdf(items).save(buf)
    return buf.getvalue()


def _one_to_pdf(inp, size="A4") -> bytes:
    ext = inp.ext
    if ext == "pdf":
        return inp.data
    if ext in OFFICE:
        return soffice_convert(inp.data, inp.name, "pdf")
    if ext == "svg":
        return _svg_pdf(inp.data, inp.name)
    if ext in IMAGES:
        buf = io.BytesIO()
        _images_pdf(_image_frames(inp.data, ext)).save(buf)
        return buf.getvalue()
    if ext == "epub":
        return _epub_pdf(inp.data, size)
    if ext == "cbz":
        return _cbz_pdf(inp.data)
    if ext in ("md", "markdown"):
        return html_to_pdf(markdown_to_html(inp.data.decode("utf-8", "replace")), size=size)
    if ext in ("html", "htm", "xhtml"):
        return html_to_pdf(inp.data.decode("utf-8", "replace"), size=size)
    if ext in ("txt", "text", "log", "json", "xml", "py", "js", "ts", "") or inp.mime.startswith("text/"):
        txt = inp.data.decode("utf-8", "replace")
        return html_to_pdf(f"<html><body><pre>{_html.escape(txt)}</pre></body></html>", size=size)
    raise EditError(f"don't know how to convert .{ext} files")


@tool("convert_to_pdf", "Convert to PDF", TO,
      "Convert files to PDF: Word, PowerPoint, Excel, Publisher, OpenDocument, "
      "RTF, images, text, Markdown, HTML, EPUB, CBZ comics. Each file becomes a PDF "
      "(or all are merged into one).",
      inputs=FILES,
      aliases=["word_to_pdf", "docx_to_pdf", "doc_to_pdf", "powerpoint_to_pdf", "pptx_to_pdf",
               "ppt_to_pdf", "excel_to_pdf", "xlsx_to_pdf", "xls_to_pdf", "publisher_to_pdf",
               "pub_to_pdf", "odt_to_pdf", "ods_to_pdf", "odp_to_pdf", "odg_to_pdf", "rtf_to_pdf",
               "text_to_pdf", "txt_to_pdf", "markdown_to_pdf", "md_to_pdf", "html_to_pdf",
               "epub_to_pdf", "svg_to_pdf", "view_as_pdf", "cbz_to_pdf"],
      options={"merge": {"type": "boolean", "default": False, "description": "Merge all results into one PDF."},
               "page_size": {"type": "string", "enum": PAPER_ENUM, "default": "A4",
                             "description": "For text/Markdown/HTML/EPUB input."}})
def convert_to_pdf(ctx, merge=False, page_size="A4"):
    results = [(inp.stem, _one_to_pdf(inp, page_size)) for inp in ctx.inputs]
    if merge and len(results) > 1:
        out = pikepdf.new()
        for _, data in results:
            out.add_pages_from(open_bytes(data))
        ctx.add_pdf(out, "converted.pdf")
    else:
        for stem, data in results:
            ctx.add_pdf(data, f"{stem}.pdf")
    ctx.result = {"converted": len(results)}


@tool("webpage_to_pdf", "Webpage to PDF", TO,
      "Save a web page (URL) as a PDF, as Chrome would print it.",
      inputs=NONE, required=["url"], requires=["chrome"], aliases=["url_to_pdf", "website_to_pdf"],
      options={"url": {"type": "string", "description": "https://..."},
               "output_name": {"type": "string"}})
def webpage_to_pdf(ctx, url, output_name=None):
    if not url.startswith(("http://", "https://")):
        if "://" in url:
            raise EditError("only http(s) URLs can be printed")
        url = "https://" + url
    data = chrome_print(url)
    host = url.split("//", 1)[-1].split("/", 1)[0] or "webpage"
    ctx.add_pdf(data, output_name or host)
    ctx.result = {"pages": len(open_bytes(data).pages)}


# -- from PDF ---------------------------------------------------------------------------
def _num(c):
    if c is None:
        return None
    s = str(c).strip()
    t = s.replace(",", "").replace("$", "").replace("€", "").replace("£", "")
    try:
        return float(t) if "." in t else int(t)
    except ValueError:
        return s


def _to_xlsx(inp, eng, pages) -> tuple[bytes, int]:
    import pdfplumber
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    n = 0
    with pdfplumber.open(io.BytesIO(inp.data)) as pl:
        for p in pages:
            for ti, rows in enumerate(pl.pages[p].extract_tables()):
                if not rows:
                    continue
                ws = wb.create_sheet(f"p{p + 1}-t{ti + 1}"[:31])
                for row in rows:
                    ws.append([_num(c) for c in row])
                n += 1
    if n == 0:  # no ruled tables: one sheet of text lines, columns split on wide gaps
        ws = wb.create_sheet("text")
        for p in pages:
            for line in eng._lines(p):
                line = sorted(line, key=lambda s: s.origin[0])
                ws.append([_num(s.text) for s in line if s.text.strip()] or [""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), n


def _to_pptx(eng, pages) -> bytes:
    from pptx import Presentation
    from pptx.util import Emu
    from .render import display_size
    prs = Presentation()
    w, h = display_size(eng, pages[0])
    prs.slide_width, prs.slide_height = Emu(int(w * 12700)), Emu(int(h * 12700))
    blank = prs.slide_layouts[6]
    for p in pages:
        slide = prs.slides.add_slide(blank)
        buf = io.BytesIO()
        eng.render_image(p, 150 / 72).save(buf, "PNG")
        buf.seek(0)
        slide.shapes.add_picture(buf, 0, 0, width=prs.slide_width, height=prs.slide_height)
        slide.notes_slide.notes_text_frame.text = eng.page_text(p)[:5000]
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _to_html(eng, pages, title) -> str:
    from .organize import to_markdown
    import markdown
    parts = [f"<section class='page' id='page-{p + 1}'>"
             + markdown.markdown(to_markdown(eng, [p]), extensions=["tables", "sane_lists"]) + "</section>"
             for p in pages]
    return ("<!doctype html><html><head><meta charset='utf-8'><title>" + _html.escape(title) +
            "</title><style>body{font-family:sans-serif;max-width:900px;margin:auto;padding:24px}"
            ".page{border-bottom:1px solid #ddd;padding:24px 0}img{max-width:100%}</style></head><body>"
            + "".join(parts) + "</body></html>")


def _pandoc(md: str, fmt: str, title: str) -> bytes:
    exe = find_program("pandoc")
    if not exe:
        raise EditError("pandoc is needed for EPUB output (brew install pandoc)")
    with tempfile.TemporaryDirectory() as tmp:
        out = os.path.join(tmp, f"out.{fmt}")
        r = subprocess.run([exe, "-f", "markdown", "-o", out, "--metadata", f"title={title}"],
                           input=md.encode(), capture_output=True, timeout=180)
        if r.returncode != 0:
            raise EditError(f"pandoc failed: {r.stderr.decode(errors='ignore')[-300:]}")
        with open(out, "rb") as fh:
            return fh.read()


IMAGE_FORMATS = {"png", "jpg", "tiff", "webp"}


def _pdf_to(ctx, target: str, pages="all", dpi=150):
    from .organize import to_markdown
    inp = ctx.inputs[0]
    ctx.pdf()  # readable / not locked
    eng = inp.engine()
    sel = parse_pages(pages, eng.page_count)
    stem = inp.stem
    if target in IMAGE_FORMATS:
        files = []
        for p in sel:
            im = eng.render_image(p, dpi / 72)
            buf = io.BytesIO()
            fmt = {"png": "PNG", "jpg": "JPEG", "tiff": "TIFF", "webp": "WEBP"}[target]
            im.save(buf, fmt, **({"quality": 90} if target in ("jpg", "webp") else {}),
                    **({"dpi": (dpi, dpi)} if target in ("png", "jpg", "tiff") else {}))
            files.append((f"{stem}-page-{p + 1}.{target}", buf.getvalue()))
        ctx.add_bundle(files, f"{stem}-{target}.zip")
        ctx.result = {"images": len(files)}
        return
    if target in ("xlsx", "ods"):
        xlsx, n = _to_xlsx(inp, eng, sel)
        ctx.result["tables"] = n
        ctx.add_file(xlsx if target == "xlsx" else soffice_convert(xlsx, f"{stem}.xlsx", "ods"),
                     f"{stem}.{target}")
    elif target in ("pptx", "odp"):
        pptx = _to_pptx(eng, sel)
        ctx.add_file(pptx if target == "pptx" else soffice_convert(pptx, f"{stem}.pptx", "odp"),
                     f"{stem}.{target}")
    elif target == "epub":
        ctx.add_file(_pandoc(to_markdown(eng, sel), "epub", stem), f"{stem}.epub", "application/epub+zip")
    elif target == "txt":
        ctx.add_file("\n\n".join(eng.page_text(p) for p in sel).encode(), f"{stem}.txt", "text/plain")
    elif target == "md":
        ctx.add_file(to_markdown(eng, sel).encode(), f"{stem}.md", "text/markdown")
    elif target == "html":
        ctx.add_file(_to_html(eng, sel, stem).encode(), f"{stem}.html", "text/html")
    else:
        raise EditError(f"unknown target {target}")
    ctx.result.setdefault("pages", len(sel))


TARGETS = {
    "xlsx": ("PDF to Excel", ["pdf_to_xls"], []),
    "pptx": ("PDF to PowerPoint", ["pdf_to_ppt"], []),
    "ods": ("PDF to ODS", [], ["soffice"]),
    "odp": ("PDF to ODP", [], ["soffice"]),
    "epub": ("PDF to EPUB", [], ["pandoc"]),
    "html": ("PDF to HTML", [], []),
    "md": ("PDF to Markdown", ["pdf_to_markdown"], []),
    "txt": ("PDF to Text", ["pdf_to_text"], []),
    "png": ("PDF to PNG", ["pdf_to_images", "pdf_to_image"], []),
    "jpg": ("PDF to JPG", ["pdf_to_jpeg"], []),
    "tiff": ("PDF to TIFF", [], []),
    "webp": ("PDF to WEBP", [], []),
}
_NAMES = {"xlsx": "excel", "pptx": "powerpoint"}


def _register(fmt: str, title: str, aliases: list, requires: list):
    opts = {"pages": PAGES_OPT}
    if fmt in IMAGE_FORMATS:
        opts["dpi"] = {"type": "integer", "default": 150}
    desc = {"xlsx": "Tables found in the PDF, one sheet each (text lines if there are no tables).",
            "pptx": "One slide per page (page image; page text in the speaker notes)."}.get(
        fmt, f"Convert the PDF to {title.split()[-1]}.")

    @tool(f"pdf_to_{_NAMES.get(fmt, fmt)}", title, FROM, desc, inputs=PDF, options=opts,
          requires=requires, aliases=aliases)
    def _fn(ctx, pages="all", dpi=150, _fmt=fmt):
        _pdf_to(ctx, _fmt, pages, dpi)
    return _fn


for _fmt, (_title, _aliases, _req) in TARGETS.items():
    _register(_fmt, _title, _aliases, _req)
